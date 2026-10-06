"""
Vocalis - Voice & Text -> Document Converter
Run:  python app.py   then open http://localhost:5000
"""
import io
import os
import re
import html
import glob
from datetime import datetime

from flask import Flask, render_template, request, send_file, jsonify

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from fpdf import FPDF

app = Flask(__name__)

FORMATS = {
    "pdf":  ("application/pdf", "pdf"),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
    "txt":  ("text/plain; charset=utf-8", "txt"),
    "md":   ("text/markdown; charset=utf-8", "md"),
    "html": ("text/html; charset=utf-8", "html"),
    "rtf":  ("application/rtf", "rtf"),
}

ALIGN_DOCX = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
ALIGN_PDF = {"left": "L", "center": "C", "right": "R", "justify": "J"}


def safe_filename(name: str) -> str:
    name = re.sub(r"[^\w\- ]+", "", name or "").strip().replace(" ", "_")
    return name or f"document_{datetime.now():%Y%m%d_%H%M%S}"


def paragraphs(text: str):
    return [p.strip() for p in re.split(r"\n\s*\n", text.strip())] if text.strip() else []


# ---------------------------------------------------------------- builders
def build_txt(title, text, author, **_):
    head = f"{title}\n{'=' * len(title)}\n"
    if author:
        head += f"By {author}  |  {datetime.now():%d %b %Y}\n"
    return (head + "\n" + text.strip() + "\n").encode("utf-8")


def build_md(title, text, author, **_):
    out = f"# {title}\n\n"
    if author:
        out += f"*By {author} — {datetime.now():%d %b %Y}*\n\n"
    out += "\n\n".join(paragraphs(text)) + "\n"
    return out.encode("utf-8")


def build_html(title, text, author, font_size, align, **_):
    paras = "\n".join(
        f"<p>{html.escape(p).replace(chr(10), '<br>')}</p>" for p in paragraphs(text)
    )
    meta = f'<p class="meta">By {html.escape(author)} · {datetime.now():%d %b %Y}</p>' if author else ""
    doc = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(title)}</title>
<style>
body{{font-family:Georgia,'Noto Serif',serif;max-width:760px;margin:48px auto;padding:0 20px;
color:#1f2330;line-height:1.7;font-size:{font_size}pt;text-align:{align}}}
h1{{font-family:system-ui,sans-serif;color:#111111;text-align:left}}
.meta{{color:#888;font-size:.85em;text-align:left;margin-top:-10px}}
</style></head><body>
<h1>{html.escape(title)}</h1>{meta}
{paras}
</body></html>"""
    return doc.encode("utf-8")


def build_docx(title, text, author, font_size, align, **_):
    d = Document()
    style = d.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(font_size)

    h = d.add_heading(title, level=0)
    for r in h.runs:
        r.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    if author:
        m = d.add_paragraph(f"By {author}  •  {datetime.now():%d %b %Y}")
        m.runs[0].italic = True
        m.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    for p in paragraphs(text):
        para = d.add_paragraph(p)
        para.alignment = ALIGN_DOCX.get(align, WD_ALIGN_PARAGRAPH.LEFT)

    d.core_properties.title = title
    if author:
        d.core_properties.author = author
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def find_unicode_font():
    """Look for a TTF that covers Indian scripts / broad Unicode."""
    here = os.path.dirname(os.path.abspath(__file__))
    candidates = glob.glob(os.path.join(here, "fonts", "*.ttf")) + [
        r"C:\Windows\Fonts\Nirmala.ttf",       # Windows: Telugu, Hindi, Tamil...
        r"C:\Windows\Fonts\NirmalaUI.ttf",
        r"C:\Windows\Fonts\arial.ttf",
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for c in candidates:
        if os.path.isfile(c):
            return c
    return None


def build_pdf(title, text, author, font_size, align, **_):
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_margins(20, 20, 20)

    font_path = find_unicode_font()
    if font_path:
        pdf.add_font("Body", "", font_path)
        family = "Body"
        try:
            pdf.set_text_shaping(True)   # proper rendering of complex scripts (needs uharfbuzz)
        except Exception:
            pass
    else:
        family = "Helvetica"
        to_latin = lambda s: s.encode("latin-1", "replace").decode("latin-1")
        title, text, author = to_latin(title), to_latin(text), to_latin(author or "")

    pdf.set_font(family, size=22)
    pdf.set_text_color(17, 17, 17)
    pdf.multi_cell(0, 11, title, new_x="LMARGIN", new_y="NEXT")

    if author:
        pdf.set_font(family, size=10)
        pdf.set_text_color(136, 136, 136)
        pdf.cell(0, 7, f"By {author}  |  {datetime.now():%d %b %Y}", new_x="LMARGIN", new_y="NEXT")

    pdf.set_draw_color(17, 17, 17)
    pdf.line(20, pdf.get_y() + 2, 190, pdf.get_y() + 2)
    pdf.ln(8)

    pdf.set_font(family, size=font_size)
    pdf.set_text_color(31, 35, 48)
    for p in paragraphs(text):
        pdf.multi_cell(0, font_size * 0.55, p, align=ALIGN_PDF.get(align, "L"),
                       new_x="LMARGIN", new_y="NEXT")
        pdf.ln(font_size * 0.35)
    return bytes(pdf.output())


def rtf_escape(s):
    out = []
    for ch in s:
        if ch in "\\{}":
            out.append("\\" + ch)
        elif ord(ch) > 127:
            code = ord(ch)
            out.append(f"\\u{code if code < 32768 else code - 65536}?")
        else:
            out.append(ch)
    return "".join(out)


def build_rtf(title, text, author, font_size, align, **_):
    a = {"left": "\\ql", "center": "\\qc", "right": "\\qr", "justify": "\\qj"}.get(align, "\\ql")
    line_break = "\\line "
    body = "".join(
        "{\\pard" + a + "\\sa200\\fs" + str(font_size * 2) + " "
        + rtf_escape(p).replace("\n", line_break) + "\\par}\n"
        for p in paragraphs(text)
    )
    meta = f"{{\\pard\\i\\fs20\\cf2 By {rtf_escape(author)}\\par}}\n" if author else ""
    rtf = (
        "{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Calibri;}}"
        "{\\colortbl;\\red17\\green17\\blue17;\\red136\\green136\\blue136;}\n"
        f"{{\\pard\\b\\fs44\\cf1 {rtf_escape(title)}\\par}}\n{meta}{body}}}"
    )
    return rtf.encode("ascii")


BUILDERS = {"pdf": build_pdf, "docx": build_docx, "txt": build_txt,
            "md": build_md, "html": build_html, "rtf": build_rtf}


# ---------------------------------------------------------------- routes
@app.route("/")
def index():
    return render_template("index.html")


@app.route("/convert", methods=["POST"])
def convert():
    data = request.get_json(silent=True) or {}
    text = (data.get("text") or "").strip()
    fmt = (data.get("format") or "pdf").lower()
    if not text:
        return jsonify(error="There's no text to convert yet."), 400
    if fmt not in BUILDERS:
        return jsonify(error=f"Unsupported format: {fmt}"), 400

    title = (data.get("title") or "Untitled Document").strip()
    try:
        font_size = max(8, min(28, int(data.get("fontSize", 12))))
    except (TypeError, ValueError):
        font_size = 12

    try:
        content = BUILDERS[fmt](
            title=title, text=text, author=(data.get("author") or "").strip(),
            font_size=font_size, align=data.get("align", "left"),
        )
    except Exception as e:  # surface a readable error to the UI
        return jsonify(error=f"Could not build {fmt.upper()}: {e}"), 500

    mime, ext = FORMATS[fmt]
    return send_file(io.BytesIO(content), mimetype=mime, as_attachment=True,
                     download_name=f"{safe_filename(title)}.{ext}")


if __name__ == "__main__":
    print("\n  Vocalis running ->  http://localhost:5000\n")
    app.run(host="127.0.0.1", port=5000, debug=True)
